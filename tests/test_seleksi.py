"""Test seleksi data: kosong, akun resmi/media, duplikat."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from aplikasi_web import main
from persiapan_data.seleksi import ALASAN_AKUN, ALASAN_DUPLIKAT, ALASAN_KOSONG, cari_kolom_akun, seleksi_data

BARIS = [
    {"user": "warga1", "text": "Bus telat lagi 😡"},
    {"user": "PT_Transjakarta", "text": "Hai Kak, mohon maaf atas ketidaknyamanannya"},
    {"user": "warga2", "text": "bus telat lagi"},            # duplikat baris 0 setelah dibersihkan
    {"user": "warga3", "text": "🙃 https://t.co/x"},         # kosong setelah dibersihkan
    {"user": "@KompasCom", "text": "Berita macet hari ini"},  # akun media, huruf & @ diabaikan
    {"user": "warga4", "text": "Parkir liar di Tanah Abang"},
]


def test_seleksi_data():
    """Setiap baris disaring dengan alasan yang tepat; yang pertama dari duplikat dipertahankan."""
    lolos, disaring = seleksi_data(BARIS, "text", "user")
    assert lolos == [0, 5]
    assert disaring == {1: ALASAN_AKUN, 2: ALASAN_DUPLIKAT, 3: ALASAN_KOSONG, 4: ALASAN_AKUN}


def test_tanpa_kolom_akun():
    """Tanpa kolom akun, hanya kosong & duplikat yang disaring."""
    lolos, disaring = seleksi_data(BARIS, "text", None)
    assert lolos == [0, 1, 4, 5]
    assert set(disaring.values()) == {ALASAN_DUPLIKAT, ALASAN_KOSONG}


def test_cari_kolom_akun():
    """Kolom akun dikenali dari nama umum, tanpa peduli huruf besar/kecil."""
    assert cari_kolom_akun(["id", "User", "text"]) == "User"
    assert cari_kolom_akun(["id", "text"]) is None


def _jalankan(seleksi: bool):
    """Upload CSV contoh, klasifikasi, lalu kembalikan (json classify, workbook export)."""
    main.SESI.clear()
    klien = TestClient(main.app)
    isi = "user,text\n" + "\n".join(f"{b['user']},{b['text']}" for b in BARIS) + "\n"
    id_sesi = klien.post("/upload", files={"file": ("d.csv", io.BytesIO(isi.encode()), "text/csv")}).json()["id"]
    data = klien.post("/classify", json={"id": id_sesi, "kolom_keluhan": "text", "seleksi": seleksi}).json()
    buku = load_workbook(io.BytesIO(klien.get(f"/export/{id_sesi}").content))
    return data, buku


def test_classify_dengan_seleksi():
    """Hanya data lolos yang dianalisis; sisanya ada di sheet Data Disaring."""
    data, buku = _jalankan(seleksi=True)
    assert len(data["hasil"]) == 2
    assert data["ringkasan_seleksi"] == {
        "Data awal": 6, ALASAN_AKUN: 2, ALASAN_DUPLIKAT: 1, ALASAN_KOSONG: 1, "Dianalisis": 2,
    }
    sheet = buku["Data Disaring"]
    assert [sel.value for sel in sheet[1]] == ["Alasan Disaring", "user", "text"]
    assert sheet.max_row == 5  # header + 4 baris disaring


def test_classify_tanpa_seleksi():
    """Seleksi bisa dimatikan: semua baris dianalisis, tanpa sheet Data Disaring."""
    data, buku = _jalankan(seleksi=False)
    assert len(data["hasil"]) == 6
    assert "Data Disaring" not in buku.sheetnames
