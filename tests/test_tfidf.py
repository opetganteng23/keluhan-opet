"""Test metode TF-IDF (Naive Bayes & SVM) dengan data tiruan di folder sementara."""

import json

import pytest
from openpyxl import Workbook, load_workbook

from metode_3_tfidf_naive_bayes.naive_bayes import buat_klasifier as nb
from metode_4_tfidf_svm.svm import buat_klasifier as svm
from persiapan_data import tfidf_umum

CONTOH = {
    ("Angkutan Umum", "Negatif"): "bus transjakarta telat lagi halte penuh sesak",
    ("Lahan Parkir", "Negatif"): "parkir liar jukir minta bayar mahal",
    ("Manajemen Lalu Lintas", "Negatif"): "macet parah lampu merah mati",
    ("Tidak Relevan", "Netral"): "promo diskon sepatu cek bio",
}


@pytest.fixture
def lingkungan(tmp_path, monkeypatch):
    """Data latih tiruan (JSONL) + file sampel uji tiruan."""
    progres = tmp_path / "progres.jsonl"
    with progres.open("w", encoding="utf-8") as f:
        for i in range(20):
            for (k, s), t in CONTOH.items():
                f.write(json.dumps({"teks_bersih": f"{t} {i}x", "Kategori": k, "Sentimen": s, "Metode": "LLM"}) + "\n")
        # Baris yang sama dengan sampel uji: harus dikeluarkan dari data latih
        f.write(json.dumps({"teks_bersih": "macet di cawang", "Kategori": "Angkutan Umum",
                            "Sentimen": "Positif", "Metode": "LLM"}) + "\n")
        # Fallback: harus diabaikan
        f.write(json.dumps({"teks_bersih": "apa saja", "Kategori": "Lahan Parkir",
                            "Sentimen": "Positif", "Metode": "Kata Kunci (fallback)"}) + "\n")
    uji = tmp_path / "hasil_engine_sampel.xlsx"
    buku = Workbook()
    sheet = buku.active
    sheet.title = "Label Engine"
    sheet.append(["ID Sampel", "Teks", "Kategori LLM"])
    sheet.append(["S001", "macet di cawang", "Manajemen Lalu Lintas"])
    sheet.append(["S002", "jukir parkir liar", "Lahan Parkir"])
    buku.save(uji)
    monkeypatch.setattr(tfidf_umum, "FILE_PROGRES", progres)
    monkeypatch.setattr(tfidf_umum, "FILE_SAMPEL_UJI", uji)
    return tmp_path


def test_data_latih_tanpa_sampel_uji_dan_fallback(lingkungan):
    """Sampel uji dan baris fallback tidak boleh masuk data latih."""
    teks, label, dikeluarkan = tfidf_umum.muat_data_latih()
    assert dikeluarkan == 1
    assert "macet di cawang" not in teks and "apa saja" not in teks
    assert len(teks) == 80 and len(label["Kategori"]) == 80


@pytest.mark.parametrize("buat, singkat", [(nb, "NB"), (svm, "SVM")])
def test_jalankan_menulis_prediksi_bukti_model(lingkungan, buat, singkat):
    """Model tersimpan, bukti dibuat, dan prediksi masuk ke kolom baru sampel uji."""
    folder = lingkungan / singkat
    folder.mkdir()
    tfidf_umum.jalankan(singkat, singkat, buat, folder)
    assert (folder / "model" / "kategori.joblib").exists()
    assert (lingkungan / f"bukti_{singkat}.xlsx").exists()
    sheet = load_workbook(lingkungan / "hasil_engine_sampel.xlsx")["Label Engine"]
    judul = [c.value for c in sheet[1]]
    k = judul.index(f"Kategori {singkat}") + 1
    assert sheet.cell(row=2, column=k).value == "Manajemen Lalu Lintas"
    assert sheet.cell(row=3, column=k).value == "Lahan Parkir"
    # Jalankan ulang: kolom diganti, bukan ditambah lagi
    tfidf_umum.jalankan(singkat, singkat, buat, folder)
    judul2 = [c.value for c in load_workbook(lingkungan / "hasil_engine_sampel.xlsx")["Label Engine"][1]]
    assert judul2.count(f"Kategori {singkat}") == 1
