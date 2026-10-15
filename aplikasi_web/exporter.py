"""Export hasil klasifikasi keluhan ke file Excel (.xlsx) menggunakan openpyxl."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet


def _tulis_header(sheet: Worksheet, header: list[str]) -> None:
    """Menulis baris header dengan format tebal pada sheet yang diberikan."""
    for kolom_ke, judul in enumerate(header, start=1):
        sel = sheet.cell(row=1, column=kolom_ke, value=judul)
        sel.font = Font(bold=True)


def _buat_sheet_hasil(workbook: Workbook, kolom_asli: list[str], baris_hasil: list[dict]) -> None:
    """Membuat sheet 'Hasil Klasifikasi' berisi kolom asli CSV + kolom hasil analisis."""
    sheet = workbook.active
    sheet.title = "Hasil Klasifikasi"

    header = kolom_asli + ["Isi Keluhan", "Teks Bersih", "Kategori", "Sentimen", "Metode"]
    _tulis_header(sheet, header)

    for baris_ke, baris in enumerate(baris_hasil, start=2):
        for kolom_ke, judul in enumerate(header, start=1):
            sheet.cell(row=baris_ke, column=kolom_ke, value=baris.get(judul, ""))


def _buat_sheet_ringkasan(workbook: Workbook, baris_hasil: list[dict]) -> None:
    """Membuat sheet 'Ringkasan': jumlah dan persentase per kategori, sentimen, dan metode."""
    sheet = workbook.create_sheet("Ringkasan")
    total = len(baris_hasil) or 1

    baris_ke = 1
    for kolom in ("Kategori", "Sentimen", "Metode"):
        jumlah_per_nilai: dict[str, int] = {}
        for baris in baris_hasil:
            nilai = baris.get(kolom, "")
            jumlah_per_nilai[nilai] = jumlah_per_nilai.get(nilai, 0) + 1

        sheet.cell(row=baris_ke, column=1, value=f"Ringkasan per {kolom}").font = Font(bold=True)
        baris_ke += 1
        for kolom_ke, judul in enumerate([kolom, "Jumlah", "Persentase"], start=1):
            sheet.cell(row=baris_ke, column=kolom_ke, value=judul).font = Font(bold=True)
        for nilai, jumlah in sorted(jumlah_per_nilai.items()):
            baris_ke += 1
            sheet.cell(row=baris_ke, column=1, value=nilai)
            sheet.cell(row=baris_ke, column=2, value=jumlah)
            sheet.cell(row=baris_ke, column=3, value=f"{jumlah / total * 100:.1f}%")
        baris_ke += 2

    sheet.cell(row=baris_ke, column=1, value="Total Data").font = Font(bold=True)
    sheet.cell(row=baris_ke, column=2, value=len(baris_hasil))


def _tulis_seleksi(workbook: Workbook, ringkasan_seleksi: dict[str, int]) -> None:
    """Menambahkan tabel 'Seleksi Data' (data awal -> alasan disaring -> dianalisis) di Ringkasan."""
    sheet = workbook["Ringkasan"]
    baris_ke = sheet.max_row + 2
    sheet.cell(row=baris_ke, column=1, value="Seleksi Data").font = Font(bold=True)
    baris_ke += 1
    for kolom_ke, judul in enumerate(["Tahap", "Jumlah"], start=1):
        sheet.cell(row=baris_ke, column=kolom_ke, value=judul).font = Font(bold=True)
    for tahap, jumlah in ringkasan_seleksi.items():
        baris_ke += 1
        disaring = tahap not in ("Data awal", "Dianalisis")
        sheet.cell(row=baris_ke, column=1, value=f"- {tahap}" if disaring else tahap)
        sheet.cell(row=baris_ke, column=2, value=-jumlah if disaring else jumlah)


def _buat_sheet_disaring(workbook: Workbook, kolom_asli: list[str], disaring: list[dict]) -> None:
    """Membuat sheet 'Data Disaring': baris yang tidak dianalisis beserta alasannya."""
    sheet = workbook.create_sheet("Data Disaring")
    header = ["Alasan Disaring"] + kolom_asli
    _tulis_header(sheet, header)
    for baris_ke, baris in enumerate(disaring, start=2):
        for kolom_ke, judul in enumerate(header, start=1):
            sheet.cell(row=baris_ke, column=kolom_ke, value=baris.get(judul, ""))


def buat_file_excel(
    kolom_asli: list[str],
    baris_hasil: list[dict],
    path_tujuan: Path,
    disaring: list[dict] | None = None,
    ringkasan_seleksi: dict[str, int] | None = None,
) -> None:
    """Membuat Excel (Hasil Klasifikasi + Ringkasan [+ Data Disaring]) dan menyimpannya ke path_tujuan."""
    workbook = Workbook()
    _buat_sheet_hasil(workbook, kolom_asli, baris_hasil)
    _buat_sheet_ringkasan(workbook, baris_hasil)
    if ringkasan_seleksi:
        _tulis_seleksi(workbook, ringkasan_seleksi)
    if disaring:
        _buat_sheet_disaring(workbook, kolom_asli, disaring)
    workbook.save(path_tujuan)
