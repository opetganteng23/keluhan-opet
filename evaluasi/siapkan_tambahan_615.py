"""Perbesar sampel uji 385 -> 1.000: ambil 615 tweet acak baru untuk dilabeli manual (definisi v2).

- Sumber: data bersih final (19.016 tweet), TIDAK termasuk 385 sampel uji lama dan 100 tweet kalibrasi.
- Lembar: Downloads/Hasil_Penelitian/_file_kerja/evaluasi_akurasi/lembar_label_v2_tambahan615.xlsx
  (format sama dengan lembar_label_v2_pelabel1.xlsx; ID Sampel T001..T615)
Pemakaian:  python evaluasi/siapkan_tambahan_615.py
"""

import json
import logging
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from persiapan_data.ekspor_siap_analisis import _aman, muat_data
from persiapan_data.pembersih import bersihkan_teks
from siapkan_kalibrasi import EMOSI, KATEGORI, PEDOMAN, SENTIMEN

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger("tambahan")

KERJA = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja"
FOLDER = KERJA / "evaluasi_akurasi"
JUMLAH = 615
BENIH = 2028
TEBAL = Font(bold=True)
ISIAN = PatternFill("solid", fgColor="FFF4CC")


def teks_terpakai() -> set[str]:
    """Teks 385 sampel uji lama + 100 kalibrasi (dikecualikan)."""
    lama = {b[2] for b in list(load_workbook(FOLDER / "lembar_label_v2_pelabel1.xlsx", read_only=True)["Pelabelan"]
                               .iter_rows(values_only=True))[1:]}
    kalib = {k["teks_bersih"] for k in json.loads((KERJA / "kalibrasi" / "kunci_kalibrasi.json").read_text(encoding="utf-8"))}
    return lama | kalib


def main() -> None:
    """Ambil 615 tweet baru, tulis lembar pelabelan + kunci id tweet."""
    baris, lolos, _ = muat_data()
    terpakai = teks_terpakai()
    kandidat = []
    for i in lolos:
        t = _aman(bersihkan_teks(baris[i]["text"]))
        if t and t not in terpakai:
            kandidat.append((i, t))
    random.seed(BENIH)
    pilih = random.sample(kandidat, JUMLAH)

    buku = Workbook()
    sheet = buku.active
    sheet.title = "Pelabelan"
    sheet.append(["No", "ID Sampel", "Teks", "Kategori", "Sentimen", "Keluhan", "Sebab", "Akibat", "Emosi", "Catatan"])
    for sel in sheet[1]:
        sel.font = TEBAL
    kunci = []
    for n, (i, t) in enumerate(pilih, 1):
        sid = f"T{n:03d}"
        sheet.append([n, sid, t, "", "", "", "", "", "", ""])
        sheet.cell(row=n + 1, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        for kolom in range(4, 10):
            sheet.cell(row=n + 1, column=kolom).fill = ISIAN
        kunci.append({"ID Sampel": sid, "id_tweet": str(baris[i]["id"]), "teks_bersih": t})
    akhir = JUMLAH + 1
    for kolom, pilihan in (("D", KATEGORI), ("E", SENTIMEN), ("F", ["Ya", "Tidak"]), ("I", EMOSI)):
        dv = DataValidation(type="list", formula1='"' + ",".join(pilihan) + '"', allow_blank=True)
        sheet.add_data_validation(dv)
        dv.add(f"{kolom}2:{kolom}{akhir}")
    for h, w in zip("ABCDEFGHIJ", [5, 9, 70, 22, 11, 9, 28, 28, 15, 22]):
        sheet.column_dimensions[h].width = w
    sheet.freeze_panes = "D2"
    pedoman = buku.create_sheet("Pedoman")
    for b in PEDOMAN:
        pedoman.append(b)
    for b in pedoman.iter_rows():
        if b[0].value and (str(b[0].value).isupper() or str(b[0].value).startswith("Pedoman")):
            b[0].font = TEBAL
        for sel in b:
            sel.alignment = Alignment(wrap_text=True, vertical="top")
    pedoman.column_dimensions["A"].width = 24
    pedoman.column_dimensions["B"].width = 110

    tujuan = FOLDER / "lembar_label_v2_tambahan615.xlsx"
    buku.save(tujuan)
    (FOLDER / "kunci_tambahan615.json").write_text(json.dumps(kunci, ensure_ascii=False, indent=1), encoding="utf-8")
    logger.info("Lembar: %s (%d tweet, dari %d kandidat)", tujuan, JUMLAH, len(kandidat))


if __name__ == "__main__":
    main()
