"""Evaluasi final 1.000 sampel (385 lama + 615 tambahan): gabungkan label manual dan siapkan label engine.

1. Gabung label manual -> lembar_label_v3_1000.xlsx (sheet Pelabelan, format sama dengan v2).
2. LLM: pakai ulang hasil 385 (batch uji_385_v3) + kirim 615 baru lewat Batch API (prompt v3.1),
   lalu pasca-proses R1/R2. Kata Kunci: dijalankan ulang untuk semua 1.000.
3. Tulis hasil_engine_sampel_v3.xlsx (sheet Label Engine + Bukti Kata Kunci).
NB/SVM (EVAL_VERSI=v3) dan IndoBERT (prediksi) menambahkan kolomnya sendiri setelah ini.

Pemakaian:  python evaluasi/siapkan_uji_v3.py
"""

import json
import logging
import shutil
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from metode_1_llm import analisis_batch as ab
from metode_1_llm.pasca_proses import koreksi
from siapkan_sampel import bukti_kata_kunci

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger("uji_v3")

KERJA = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja"
FOLDER = KERJA / "evaluasi_akurasi"
DIISI_PENELITI = Path.home() / "Downloads" / "Hasil_Penelitian" / "06_Evaluasi_Akurasi" / "LABELI_ULANG_615_tweet.xlsx"
TEBAL = Font(bold=True)


def gabung_label_manual() -> list[tuple]:
    """Salin lembar 615 yang diisi peneliti ke folder kerja, lalu gabung dengan 385 -> 1.000 baris."""
    shutil.copy2(DIISI_PENELITI, FOLDER / "lembar_label_v2_tambahan615.xlsx")
    baris = []
    for nama in ("lembar_label_v2_pelabel1.xlsx", "lembar_label_v2_tambahan615.xlsx"):
        semua = list(load_workbook(FOLDER / nama, read_only=True)["Pelabelan"].iter_rows(values_only=True))
        baris += [r for r in semua[1:] if r[1]]
    buku = Workbook()
    sheet = buku.active
    sheet.title = "Pelabelan"
    sheet.append(["No", "ID Sampel", "Teks", "Kategori", "Sentimen", "Keluhan", "Sebab", "Akibat", "Emosi", "Catatan"])
    for i, r in enumerate(baris, 1):
        sheet.append([i] + list(r[1:10]))
    buku.save(FOLDER / "lembar_label_v3_1000.xlsx")
    logger.info("Label manual gabungan: %d baris", len(baris))
    return baris


def label_llm(sampel_baru: list[dict]) -> dict[str, dict]:
    """Hasil LLM: 385 dari batch lama + 615 dari batch baru (dikirim jika belum)."""
    hasil = {}
    for b in map(json.loads, (KERJA / "analisis_llm" / "hasil_uji_385_v3.jsonl").open(encoding="utf-8")):
        b["Koreksi"] = koreksi(b)
        hasil[b["ID Sampel"]] = b
    nama = "uji_615_v3"
    if not ab.path_status(nama).exists():
        ab.kirim(nama, sampel_baru, "615 sampel uji tambahan (evaluasi final 1.000, prompt v3.1)")
    while ab.cek(nama) != "ended":
        time.sleep(30)
    for b in ab.ambil(nama):
        b["Koreksi"] = koreksi(b)
        hasil[b["ID Sampel"]] = b
    return hasil


def main() -> None:
    """Gabung label manual, jalankan LLM + Kata Kunci, tulis hasil_engine_sampel_v3.xlsx."""
    manual = gabung_label_manual()
    baru = [{"id": r[1], "ID Sampel": r[1], "teks_bersih": r[2] or ""} for r in manual if str(r[1]).startswith("T")]
    llm = label_llm(baru)

    buku = Workbook()
    sheet = buku.active
    sheet.title = "Label Engine"
    sheet.append(["ID Sampel", "Teks", "Kategori LLM", "Sentimen LLM", "Keluhan LLM", "Sebab LLM", "Akibat LLM",
                  "Emosi LLM", "Kategori Kata Kunci", "Sentimen Kata Kunci", "Keluhan Kata Kunci"])
    bukti = buku.create_sheet("Bukti Kata Kunci")
    bukti.append(["ID Sampel", "Teks", "Cocok: Angkutan Umum", "Cocok: Lahan Parkir", "Cocok: Manajemen Lalu Lintas",
                  "Alasan Kategori", "Kata Positif", "Kata Negatif"])
    for r in manual:
        sid, teks = r[1], r[2] or ""
        l = llm[sid]
        kk = bukti_kata_kunci(teks)
        sheet.append([sid, teks, l["Kategori"], l["Sentimen"], l["Keluhan"], l["Sebab"], l["Akibat"], l["Emosi"],
                      kk["kategori"], kk["sentimen"], "Ya" if kk["sentimen"] == "Negatif" else "Tidak"])
        bukti.append([sid, teks, kk["cocok"]["Angkutan Umum"], kk["cocok"]["Lahan Parkir"],
                      kk["cocok"]["Manajemen Lalu Lintas"], kk["alasan_kategori"], kk["kata_positif"], kk["kata_negatif"]])
    for sh in (sheet, bukti):
        for sel in sh[1]:
            sel.font = TEBAL
        sh.column_dimensions["B"].width = 60
    tujuan = FOLDER / "hasil_engine_sampel_v3.xlsx"
    buku.save(tujuan)
    logger.info("Label engine 1.000 sampel (LLM + Kata Kunci): %s", tujuan)


if __name__ == "__main__":
    main()
