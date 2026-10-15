"""Analisis LLM untuk banyak tweet sekaligus (sampel acak dari data bersih), bisa dilanjutkan.

- Sampel acak N tweet bersih (benih tetap -> bisa diulang persis).
- Diproses per bagian 500 tweet; hasil tiap bagian langsung disimpan ke file progres (JSONL),
  sehingga jika terputus, menjalankan ulang hanya memproses sisanya (tidak bayar dua kali).
- Akhirnya dibuat Excel: hasil per tweet + ringkasan kategori x sentimen + biaya.

Pemakaian:  python metode_1_llm/analisis_massal.py [jumlah]   (default 15000)
"""

import asyncio
import json
import logging
import random
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # agar bisa dijalankan langsung

from openpyxl import Workbook
from openpyxl.styles import Font

from metode_1_llm.llm_analyzer import MODEL, analisis_dengan_llm
from persiapan_data.ekspor_siap_analisis import _aman, muat_data
from persiapan_data.pembersih import bersihkan_teks

logging.basicConfig(level=logging.INFO, format="%(message)s")
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("httpx2").setLevel(logging.WARNING)
logger = logging.getLogger("analisis_massal")

BENIH_ACAK = 2026
UKURAN_BAGIAN = 500
HARGA_INPUT, HARGA_OUTPUT = 1.0, 5.0  # USD per 1 juta token (claude-haiku-4-5)
FOLDER = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja" / "analisis_llm"
TEBAL = Font(bold=True)
KATEGORI = ["Angkutan Umum", "Lahan Parkir", "Manajemen Lalu Lintas", "Tidak Relevan"]
SENTIMEN = ["Negatif", "Netral", "Positif"]


def file_progres(n: int) -> Path:
    """File JSONL berisi hasil per tweet yang sudah selesai (satu baris = satu tweet)."""
    return FOLDER / f"progres_llm_{n}.jsonl"


def muat_progres(path: Path) -> dict[str, dict]:
    """Hasil yang sudah tersimpan: {id tweet: hasil}."""
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as f:
        return {d["id"]: d for d in map(json.loads, f)}


async def proses(sampel: list[dict], path: Path) -> tuple[int, int]:
    """Proses tweet yang belum selesai per bagian; kembalikan total token input & output sesi ini."""
    selesai = muat_progres(path)
    sisa = [s for s in sampel if s["id"] not in selesai]
    logger.info("Sudah selesai %d, sisa %d tweet", len(selesai), len(sisa))
    tok_in = tok_out = 0
    for awal in range(0, len(sisa), UKURAN_BAGIAN):
        bagian = sisa[awal:awal + UKURAN_BAGIAN]
        jejak: list[dict] = []
        label = await analisis_dengan_llm([s["teks_bersih"] for s in bagian], jejak)
        with path.open("a", encoding="utf-8") as f:
            for s, l in zip(bagian, label):
                f.write(json.dumps({**s, **l}, ensure_ascii=False) + "\n")
        tok_in += sum(j["token_input"] for j in jejak)
        tok_out += sum(j["token_output"] for j in jejak)
        biaya = tok_in / 1e6 * HARGA_INPUT + tok_out / 1e6 * HARGA_OUTPUT
        logger.info("[%s] %d/%d tweet selesai | biaya sesi ini $%.2f",
                    datetime.now().strftime("%H:%M"), len(selesai) + awal + len(bagian), len(sampel), biaya)
    return tok_in, tok_out


def simpan_excel(hasil: list[dict], n: int, tok_in: int, tok_out: int) -> Path:
    """Excel: Hasil per tweet, Ringkasan kategori x sentimen, Info."""
    buku = Workbook()
    sheet = buku.active
    sheet.title = "Hasil LLM"
    sheet.append(["id", "date", "user", "text", "Teks Bersih", "seed_query", "Kategori", "Sentimen", "Metode"])
    for d in hasil:
        sheet.append([d["id"], d["date"], d["user"], _aman(d["text"]), _aman(d["teks_bersih"]), d["query"],
                      d["Kategori"], d["Sentimen"], d["Metode"]])
    for sel in sheet[1]:
        sel.font = TEBAL
    for h, w in zip("ABCDEFGHI", [21, 26, 18, 60, 60, 30, 22, 10, 10]):
        sheet.column_dimensions[h].width = w
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    ring = buku.create_sheet("Ringkasan")
    silang = Counter((d["Kategori"], d["Sentimen"]) for d in hasil)
    ring.append(["Kategori"] + SENTIMEN + ["Total", "% dari semua"])
    for k in KATEGORI:
        baris = [silang[(k, s)] for s in SENTIMEN]
        ring.append([k] + baris + [sum(baris), round(sum(baris) / len(hasil) * 100, 1)])
    total = [sum(silang[(k, s)] for k in KATEGORI) for s in SENTIMEN]
    ring.append(["Total"] + total + [len(hasil), 100.0])
    relevan = [d for d in hasil if d["Kategori"] != "Tidak Relevan"]
    keluhan = [d for d in relevan if d["Sentimen"] == "Negatif"]
    ring.append([])
    ring.append(["Tweet relevan", len(relevan)])
    ring.append(["Keluhan (relevan + negatif)", len(keluhan)])
    ring.append(["Fallback kata kunci", sum(d["Metode"] != "LLM" for d in hasil)])
    for b in (1, len(KATEGORI) + 2):
        for sel in ring[b]:
            sel.font = TEBAL
    ring.column_dimensions["A"].width = 28

    info = buku.create_sheet("Info")
    biaya = tok_in / 1e6 * HARGA_INPUT + tok_out / 1e6 * HARGA_OUTPUT
    for b in [["Dibuat", datetime.now().strftime("%d-%m-%Y %H:%M")], ["Model", MODEL],
              ["Sampel", f"{n} tweet acak dari data bersih (benih {BENIH_ACAK})"],
              ["Token sesi terakhir", f"input {tok_in}, output {tok_out} (≈ ${biaya:.2f})"]]:
        info.append(b)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70

    tujuan = FOLDER / f"hasil_llm_{n}.xlsx"
    buku.save(tujuan)
    return tujuan


def main() -> None:
    """Ambil sampel acak, proses yang belum selesai, lalu tulis Excel."""
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 15000
    FOLDER.mkdir(parents=True, exist_ok=True)
    baris, lolos, _ = muat_data()
    random.seed(BENIH_ACAK)
    pilih = random.sample(lolos, min(n, len(lolos)))
    sampel = [{"id": baris[i]["id"], "date": baris[i]["date"], "user": baris[i]["user"], "text": baris[i]["text"],
               "query": baris[i]["query"], "teks_bersih": bersihkan_teks(baris[i]["text"])} for i in pilih]
    path = file_progres(n)
    tok_in, tok_out = asyncio.run(proses(sampel, path))
    hasil = list(muat_progres(path).values())
    tujuan = simpan_excel(hasil, n, tok_in, tok_out)
    logger.info("SELESAI %d tweet | Excel: %s", len(hasil), tujuan)


if __name__ == "__main__":
    main()
