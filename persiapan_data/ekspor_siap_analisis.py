"""Membuat Excel 'data siap analisis' dari hasil crawling (crawling/tweets_raw.csv).

Pemakaian:
    python persiapan_data/ekspor_siap_analisis.py            -> buat Excel di folder Downloads
    python persiapan_data/ekspor_siap_analisis.py --hitung   -> hanya cetak jumlah data bersih
"""

import csv
import logging
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # agar bisa dijalankan langsung

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from metode_2_kata_kunci.classifier import klasifikasi_keluhan
from persiapan_data.pembersih import bersihkan_teks
from persiapan_data.seleksi import seleksi_data

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger("ekspor")

FILE_MENTAH = Path(__file__).resolve().parent.parent.parent / "crawling" / "tweets_raw.csv"
FOLDER_TUJUAN = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja" / "data_siap_analisis"
TEBAL = Font(bold=True)


def _aman(nilai) -> str:
    """Membuang karakter kontrol yang tidak diizinkan di sel Excel."""
    return ILLEGAL_CHARACTERS_RE.sub("", str(nilai)) if nilai is not None else ""


def muat_data() -> tuple[list[dict], list[int], dict[int, str]]:
    """Membaca data mentah lalu menjalankan seleksi (kolom teks 'text', akun 'user')."""
    with FILE_MENTAH.open(encoding="utf-8", newline="") as berkas:
        baris = list(csv.DictReader(berkas))
    lolos, disaring = seleksi_data(baris, "text", "user")
    return baris, lolos, disaring


def buat_excel(catatan_sumber: str = "") -> Path:
    """Membuat Excel tiga sheet (Siap Analisis, Data Disaring, Info) di folder Downloads."""
    baris, lolos, disaring = muat_data()
    sekarang = datetime.now()
    buku = Workbook()

    sheet = buku.active
    sheet.title = "Siap Analisis"
    sheet.append(["No", "id", "date", "user", "text", "Teks Bersih", "reply_to", "query"])
    for nomor, i in enumerate(lolos, 1):
        r = baris[i]
        sheet.append([nomor, _aman(r["id"]), _aman(r["date"]), _aman(r["user"]), _aman(r["text"]),
                      _aman(bersihkan_teks(r["text"])), _aman(r["reply_to"]), _aman(r["query"])])
    for sel in sheet[1]:
        sel.font = TEBAL
    for huruf, lebar in zip("ABCDEFGH", [7, 21, 26, 18, 60, 60, 16, 32]):
        sheet.column_dimensions[huruf].width = lebar
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    sheet2 = buku.create_sheet("Data Disaring")
    sheet2.append(["Alasan Disaring", "id", "date", "user", "text"])
    for i, alasan in sorted(disaring.items()):
        r = baris[i]
        sheet2.append([alasan, _aman(r["id"]), _aman(r["date"]), _aman(r["user"]), _aman(r["text"])])
    for sel in sheet2[1]:
        sel.font = TEBAL
    for huruf, lebar in zip("ABCDE", [27, 21, 26, 18, 70]):
        sheet2.column_dimensions[huruf].width = lebar
    sheet2.freeze_panes = "A2"
    sheet2.auto_filter.ref = sheet2.dimensions

    sheet3 = buku.create_sheet("Info")
    per_tahun = Counter(baris[i]["date"][:4] for i in lolos)
    kategori = Counter(klasifikasi_keluhan(bersihkan_teks(baris[i]["text"])) for i in lolos)
    isi = [["Info Data"], ["Dibuat", sekarang.strftime("%d-%m-%Y %H:%M")],
           ["Sumber", catatan_sumber or "Crawling Twitter/X"], [],
           ["Seleksi Data"], ["Tahap", "Jumlah"], ["Data mentah", len(baris)]]
    isi += [[f"- {alasan}", -jumlah] for alasan, jumlah in Counter(disaring.values()).most_common()]
    isi += [["Siap analisis", len(lolos)], [], ["Siap analisis per tahun"], ["Tahun", "Jumlah"]]
    isi += [[tahun, jumlah] for tahun, jumlah in sorted(per_tahun.items())]
    isi += [[], ["Perkiraan kategori (metode kata kunci, sementara)"], ["Kategori", "Jumlah"]]
    isi += [[nama, jumlah] for nama, jumlah in kategori.most_common()]
    isi += [[], ["Catatan"],
            ["Teks Bersih = teks tanpa emoji, link, mention, awalan RT, kode HTML (dipakai untuk analisis)."],
            ["Belum termasuk penyaringan 'Tidak Relevan' oleh LLM; jumlah final akan lebih kecil."]]
    for b in isi:
        sheet3.append(b)
    for b in sheet3.iter_rows():
        if b[0].value in ("Info Data", "Seleksi Data", "Siap analisis per tahun", "Catatan") \
                or str(b[0].value).startswith("Perkiraan"):
            b[0].font = TEBAL
    sheet3.column_dimensions["A"].width = 45
    sheet3.column_dimensions["B"].width = 70

    tujuan = FOLDER_TUJUAN / f"data_siap_analisis_{sekarang:%Y%m%d_%H%M}.xlsx"
    buku.save(tujuan)
    logger.info("FILE: %s", tujuan)
    logger.info("mentah %d | siap analisis %d | disaring %s", len(baris), len(lolos), dict(Counter(disaring.values())))
    logger.info("per tahun %s", dict(sorted(per_tahun.items())))
    logger.info("kategori (kata kunci) %s", dict(kategori))
    return tujuan


if __name__ == "__main__":
    if "--hitung" in sys.argv:
        _, lolos, _ = muat_data()
        print(len(lolos))
    else:
        buat_excel(" ".join(a for a in sys.argv[1:]))
