"""Ekspansi seed word: cari kandidat kata/frasa baru dari tweet hasil crawling.

Metode (ko-okurensi + lift):
1. Setiap tweet bersih dikelompokkan ke kategori seed word yang menemukannya.
2. Teks dipecah menjadi unigram & bigram (huruf kecil, tanpa stopword, tanpa angka).
3. Untuk setiap istilah dihitung:
   - df_kategori : jumlah tweet di kategori itu yang memuat istilah
   - df_total    : jumlah tweet di semua kategori yang memuat istilah
   - lift        : (df_kategori / n_kategori) / (df_total / n_total)
     -> > 1 berarti istilah lebih khas untuk kategori tsb dibanding rata-rata
4. Istilah dengan df_kategori >= MIN_DF dan lift >= MIN_LIFT diurutkan
   berdasarkan skor = df_kategori x log(lift).
Hasilnya kandidat; pemilihan seed word baru dilakukan manusia (kolom "Pilih").

Pemakaian:  python persiapan_data/ekspansi_seed_word.py            (semua tweet bersih, per kategori seed word)
            python persiapan_data/ekspansi_seed_word.py --keluhan  (hanya KELUHAN: relevan + negatif menurut LLM)
"""

import json
import logging
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # agar bisa dijalankan langsung

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from persiapan_data.ekspor_siap_analisis import _aman, muat_data
from persiapan_data.pembersih import bersihkan_teks

logging.basicConfig(level=logging.INFO, format="%(message)s")
logger = logging.getLogger("ekspansi")

FILE_CRAWLER = Path(__file__).resolve().parent.parent.parent / "crawling" / "crawl_per_bulan.py"
MIN_DF = 15          # minimal muncul di 15 tweet kategori itu
MIN_LIFT = 1.5       # minimal 1,5x lebih khas dibanding rata-rata
MIN_DF_KELUHAN = 8  # korpus keluhan lebih kecil -> ambang lebih rendah
TOP_N = 60           # kandidat per kategori
TEBAL = Font(bold=True)
ISIAN = PatternFill("solid", fgColor="FFF4CC")

# Kategori tiap seed word (urutan sama dengan KATA_KUNCI di crawl_per_bulan.py)
KATEGORI_SEED = ["Angkutan Umum"] * 10 + ["Lahan Parkir"] * 6 + ["Manajemen Lalu Lintas"] * 12 + ["Dishub (umum)"] * 4
SEED_LAMA = {"rambu lalu lintas jakarta": "rambu jakarta"}  # query crawl awal -> seed word

STOPWORD = set("""
yang dan di ke dari ini itu dengan untuk pada adalah dalam tidak juga akan sudah ada bisa karena atau
saya aku gue gua gw lu lo kamu kalian kita kami mereka dia ia anda nya kak min mas mbak pak bu bang bro sis
ya yaa yah ga gak nggak enggak tak tdk gk ngga engga aja saja deh dong sih kok kan lah loh nih tuh pun
lagi masih mau udah sdh blm belum baru cuma hanya jadi buat bikin sama kalo kalau klo kl tapi tp namun
apa siapa mana kapan gimana bagaimana kenapa mengapa berapa sangat banget bgt bener benar emang memang
oleh agar supaya bahwa sebagai seperti sampai hingga setelah sebelum saat ketika jam hari tadi nanti
pagi siang sore malam kemarin besok sekarang terus trus lebih paling semua setiap tiap para banyak
sedikit satu dua tiga orang ada adanya hal begitu gitu gini begini sini situ sana mungkin harus perlu
via info kok sih wkwk wkwkwk haha hahaha hehe the and for you this that are with amp rt yg dgn utk dr
krn karna jd sm aja dah kah pas biar juga lg lgi gmn knp org trs sy sya tau tahu liat lihat kasih
""".split())


def baca_seed_word() -> list[str]:
    """Membaca daftar seed word dari crawl_per_bulan.py (variabel KATA_KUNCI)."""
    teks = FILE_CRAWLER.read_text(encoding="utf-8")
    blok = re.search(r"KATA_KUNCI = \[(.*?)\]", teks, re.S).group(1)
    return re.findall(r'"([^"]+)"', blok)


def seed_dari_query(query: str) -> str:
    """'macet jakarta [2026-03]' atau 'macet jakarta -filter:retweets' -> 'macet jakarta'."""
    seed = re.sub(r"\s*\[\d{4}-\d{2}\]$", "", query or "").replace(" -filter:retweets", "").strip()
    return SEED_LAMA.get(seed, seed)


def token(teks: str) -> list[str]:
    """Kata huruf kecil, minimal 3 huruf, bukan stopword, tanpa angka."""
    return [k for k in re.findall(r"[a-z]+", teks.lower()) if len(k) >= 3 and k not in STOPWORD]


def istilah(teks: str) -> set[str]:
    """Himpunan unigram + bigram (bigram dari token yang bersebelahan setelah stopword dibuang)."""
    t = token(teks)
    return set(t) | {f"{a} {b}" for a, b in zip(t, t[1:])}


def bagian_seed(seed_words: list[str]) -> set[str]:
    """Semua kata/frasa yang sudah menjadi bagian seed word (tidak dijadikan kandidat)."""
    hasil = set()
    for s in seed_words:
        t = re.findall(r"[a-z]+", s.lower())
        hasil |= set(t) | {f"{a} {b}" for a, b in zip(t, t[1:])}
    return hasil


def hitung_kandidat() -> tuple[dict[str, list[dict]], dict[str, int], int]:
    """Kandidat istilah per kategori + jumlah tweet per kategori + total tweet."""
    seed_words = baca_seed_word()
    kategori_seed = dict(zip(seed_words, KATEGORI_SEED))
    sudah = bagian_seed(seed_words)
    baris, lolos, _ = muat_data()

    df_kat: dict[str, Counter] = defaultdict(Counter)
    df_total: Counter = Counter()
    n_kat: Counter = Counter()
    contoh: dict[tuple[str, str], str] = {}
    for i in lolos:
        kat = kategori_seed.get(seed_dari_query(baris[i]["query"]))
        if not kat:
            continue
        teks = bersihkan_teks(baris[i]["text"])
        n_kat[kat] += 1
        for s in istilah(teks):
            if s in sudah:
                continue
            df_kat[kat][s] += 1
            df_total[s] += 1
            contoh.setdefault((kat, s), teks)

    n_total = sum(n_kat.values())
    hasil = {}
    for kat, df in df_kat.items():
        daftar = []
        for s, n in df.items():
            if n < MIN_DF:
                continue
            lift = (n / n_kat[kat]) / (df_total[s] / n_total)
            if lift < MIN_LIFT:
                continue
            daftar.append({"istilah": s, "df": n, "df_total": df_total[s], "lift": lift,
                           "skor": n * math.log(lift), "contoh": contoh[(kat, s)]})
        hasil[kat] = sorted(daftar, key=lambda d: d["skor"], reverse=True)[:TOP_N]
    return hasil, dict(n_kat), n_total


FILE_LABEL_LLM = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja" / "analisis_llm" / "progres_llm_15000.jsonl"


def hitung_kandidat_keluhan() -> tuple[dict[str, list[dict]], dict[str, int], int]:
    """Kandidat dari tweet KELUHAN (kategori relevan + sentimen negatif menurut LLM).

    Kelompok target = keluhan per kategori LLM; pembanding = seluruh tweet berlabel LLM.
    lift = P(istilah | keluhan kategori) / P(istilah | semua tweet).
    """
    sudah = bagian_seed(baca_seed_word())
    df_kat: dict[str, Counter] = defaultdict(Counter)
    df_total: Counter = Counter()
    n_kat: Counter = Counter()
    contoh: dict[tuple[str, str], str] = {}
    n_total = 0
    with FILE_LABEL_LLM.open(encoding="utf-8") as f:
        for d in map(json.loads, f):
            n_total += 1
            keluhan = d["Kategori"] != "Tidak Relevan" and d["Sentimen"] == "Negatif"
            kat = d["Kategori"]
            if keluhan:
                n_kat[kat] += 1
            for s in istilah(d["teks_bersih"]):
                if s in sudah:
                    continue
                df_total[s] += 1
                if keluhan:
                    df_kat[kat][s] += 1
                    contoh.setdefault((kat, s), d["teks_bersih"])
    hasil = {}
    for kat, df in df_kat.items():
        daftar = []
        for s, n in df.items():
            if n < MIN_DF_KELUHAN:
                continue
            lift = (n / n_kat[kat]) / (df_total[s] / n_total)
            if lift < MIN_LIFT:
                continue
            daftar.append({"istilah": s, "df": n, "df_total": df_total[s], "lift": lift,
                           "skor": n * math.log(lift), "contoh": contoh[(kat, s)]})
        hasil[kat] = sorted(daftar, key=lambda d: d["skor"], reverse=True)[:TOP_N]
    return hasil, dict(n_kat), n_total


def simpan_excel(hasil: dict[str, list[dict]], n_kat: dict[str, int], n_total: int,
                 keluhan: bool = False) -> Path:
    """Excel kandidat per kategori (kolom Pilih berupa dropdown Ya/Tidak) + sheet metode."""
    buku = Workbook()
    metode = buku.active
    metode.title = "Metode"
    isi = [
        ["Ekspansi Seed Word – Kandidat Istilah Baru"],
        ["Dibuat", datetime.now().strftime("%d-%m-%Y %H:%M")],
        ["Korpus", f"KELUHAN (relevan + negatif menurut LLM) dari {n_total} tweet berlabel LLM" if keluhan
                   else f"{n_total} tweet bersih dari 32 seed word"],
        [],
        ["Langkah", ("1) Hanya tweet keluhan (kategori relevan + sentimen negatif, label LLM) per kategori; "
                     "pembanding = seluruh tweet berlabel LLM. "
                     "2) Unigram & bigram tanpa stopword. 3) df dan lift = P(istilah|keluhan kategori) / P(istilah|semua). "
                     "4) Disaring df >= %d dan lift >= %.1f, diurutkan skor = df x ln(lift)." % (MIN_DF_KELUHAN, MIN_LIFT))
                    if keluhan else
                    "1) Tweet dikelompokkan ke kategori seed word yang menemukannya. "
                    "2) Teks dibersihkan, dipecah jadi unigram & bigram, stopword dibuang. "
                    "3) Dihitung df (jumlah tweet memuat istilah) dan lift = P(istilah|kategori) / P(istilah). "
                    "4) Disaring df >= %d dan lift >= %.1f, diurutkan skor = df x ln(lift)." % (MIN_DF, MIN_LIFT)],
        ["Membaca lift", "Lift 2,0 = istilah 2x lebih sering muncul di kategori ini dibanding rata-rata semua kategori."],
        ["Tugas peneliti", "Isi kolom 'Pilih' (Ya/Tidak) pada tiap sheet kategori. Istilah 'Ya' menjadi seed word baru."],
        [],
        ["Kategori", "Jumlah tweet keluhan" if keluhan else "Jumlah tweet"],
    ] + [[k, n] for k, n in sorted(n_kat.items(), key=lambda x: -x[1])]
    for b in isi:
        metode.append(b)
    for b in metode.iter_rows():
        if b[0].value in ("Ekspansi Seed Word – Kandidat Istilah Baru", "Kategori"):
            b[0].font = TEBAL
        for sel in b:
            sel.alignment = Alignment(wrap_text=True, vertical="top")
    metode.column_dimensions["A"].width = 22
    metode.column_dimensions["B"].width = 100

    for kat in sorted(hasil, key=lambda k: -n_kat[k]):
        sheet = buku.create_sheet(kat[:31])
        sheet.append(["Peringkat", "Istilah", "Jumlah Tweet (kategori)", "Jumlah Tweet (semua)", "Lift",
                      "Pilih (Ya/Tidak)", "Contoh Tweet"])
        for sel in sheet[1]:
            sel.font = TEBAL
        for r, d in enumerate(hasil[kat], 1):
            sheet.append([r, d["istilah"], d["df"], d["df_total"], round(d["lift"], 2), "", _aman(d["contoh"])])
            sheet.cell(row=r + 1, column=6).fill = ISIAN
            sheet.cell(row=r + 1, column=7).alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"Ya,Tidak"', allow_blank=True)
        sheet.add_data_validation(dv)
        dv.add(f"F2:F{len(hasil[kat]) + 1}")
        for h, w in zip("ABCDEFG", [10, 24, 14, 14, 8, 14, 80]):
            sheet.column_dimensions[h].width = w
        sheet.freeze_panes = "C2"

    nama = "kandidat_ekspansi_seed_word_keluhan.xlsx" if keluhan else "kandidat_ekspansi_seed_word.xlsx"
    tujuan = Path.home() / "Downloads" / "Hasil_Penelitian" / "_file_kerja" / "seed_word" / nama
    buku.save(tujuan)
    return tujuan


if __name__ == "__main__":
    mode_keluhan = "--keluhan" in sys.argv
    hasil, n_kat, n_total = hitung_kandidat_keluhan() if mode_keluhan else hitung_kandidat()
    tujuan = simpan_excel(hasil, n_kat, n_total, mode_keluhan)
    logger.info("Korpus %d tweet | per kategori %s", n_total, n_kat)
    for kat, daftar in hasil.items():
        logger.info("%s: %s", kat, ", ".join(d["istilah"] for d in daftar[:15]))
    logger.info("File: %s", tujuan)
